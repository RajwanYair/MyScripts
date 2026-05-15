from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_heading(document: Document, text: str) -> None:
    heading = document.add_paragraph()
    heading.style = document.styles["Heading 2"]
    run = heading.add_run(text)
    run.font.size = Pt(13)


def add_role(document: Document, title: str, company: str, location: str, years: str, bullets: list[str]) -> None:
    header = document.add_paragraph()
    run = header.add_run(f"{title} | {company} | {location} | {years}")
    run.bold = True
    run.font.size = Pt(11.5)
    for bullet in bullets:
        para = document.add_paragraph(bullet)
        para.style = document.styles["List Bullet"]
        para_format = para.paragraph_format
        para_format.space_before = Pt(0)
        para_format.space_after = Pt(1)


def main() -> None:
    document = Document()

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Yair Rajwan")
    title_run.bold = True
    title_run.font.size = Pt(22)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("CAD & Engineering Compute Leader")
    subtitle_run.font.size = Pt(12)
    subtitle_run.bold = True

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact.add_run("Jerusalem, Israel | 054-7887-532 | Yair.rajwan@gmail.com")
    contact_run.font.size = Pt(10.5)

    add_heading(document, "Professional Summary")
    summary = ("Engineering-compute leader with two decades at Intel orchestrating simulation enablement, "
               "embedded compute, and IT operations. Known for unifying design, validation, and backend teams "
               "to deliver IEEE-compliant workflows, resilient capacity, and predictable handoffs. Combine "
               "vendor partnership management with hands-on UNIX and RTL expertise to remove friction and "
               "accelerate silicon delivery.")
    paragraph = document.add_paragraph(summary)
    paragraph.paragraph_format.space_after = Pt(6)

    add_heading(document, "Core Strengths")
    strengths = [
        "EDA flow leadership with Siemens QuestaSim, Cadence Xcelium/Specman/Genus/Conformal, and Synopsys VCS/Fusion",
        "Compute infrastructure planning, capacity management, and automation across hybrid environments",
        "Cross-functional coaching of designers, validators, and backend engineers to streamline handoffs",
        "Standards compliance champion ensuring robust LRM/IEEE adherence across diverse toolchains",
        "Continuous-improvement mindset focused on data-driven decisions and scalable documentation"
    ]
    for strength in strengths:
        para = document.add_paragraph(strength)
        para.style = document.styles["List Bullet"]
        para.paragraph_format.space_after = Pt(1)

    add_heading(document, "Experience")
    add_role(
        document,
        title="Team Lead, Simulation Enablement",
        company="Intel Corporation",
        location="Israel",
        years="Mid 2019 – Present",
        bullets=[
            "Direct two engineers and an intern delivering EDA tool onboarding, backend handoff readiness, and compute support for cross-site design teams.",
            "Own strategic relationships with Siemens, Cadence, and Synopsys to prioritize roadmap items, negotiate fixes, and resolve blockers before they impact delivery.",
            "Institutionalized IEEE/LRM compliance guardrails that shield flows from syntax drift and reduce debug loops between teams.",
            "Partner with validation and backend leaders on capacity planning, ensuring resilient infrastructure for tape-out-critical milestones."
        ]
    )

    add_role(
        document,
        title="Simulation Enablement Engineer",
        company="Intel Corporation",
        location="Israel",
        years="2018 – Mid 2019",
        bullets=[
            "Embedded with validation groups to troubleshoot tool behavior, optimize compute and storage allocation, and elevate productivity.",
            "Built reusable playbooks and knowledge bases that accelerated onboarding for new verification projects.",
            "Maintained uninterrupted execution for large-scale regressions by proactively monitoring infrastructure health."
        ]
    )

    add_role(
        document,
        title="Lead Engineer, Frontend-to-Backend Shift-Left Program",
        company="Intel Corporation",
        location="Israel",
        years="2014 – 2018",
        bullets=[
            "Led cross-functional initiative that improved RTL handoff quality and minimized backend rework across multiple SoC programs.",
            "Developed automation for RTL filelist sorting and correctness validation via synthesis backend tools, shortening feedback cycles.",
            "Championed best practices that helped design teams adopt consistent methodologies and documentation."
        ]
    )

    add_role(
        document,
        title="Embedded Compute Engineer",
        company="Intel Corporation",
        location="Israel",
        years="2010 – 2014",
        bullets=[
            "Implemented RTL and validation solutions for networking and client engineering groups, bridging software and hardware teams.",
            "Collaborated with geographically distributed teams to integrate embedded systems into production environments.",
            "Translated customer requirements into maintainable compute and firmware deliverables."
        ]
    )

    add_role(
        document,
        title="IT Systems Engineer",
        company="Intel Corporation",
        location="Israel",
        years="2005 – 2010",
        bullets=[
            "Owned engineering compute platforms with emphasis on OS maintenance, installation, and lifecycle management.",
            "Managed Intel's internal compute cloud and storage ecosystems to protect data integrity and availability.",
            "Provided responsive support to development teams, aligning infrastructure capabilities with roadmap needs."
        ]
    )

    add_role(
        document,
        title="IT Systems Specialist (Contractor)",
        company="Matrix Inc. at Intel",
        location="Israel",
        years="2001 – 2005",
        bullets=[
            "Oversaw datacenter and UNIX environments spanning AIX, HPUX, Solaris, and Linux, ensuring uptime for development workloads.",
            "Delivered infrastructure upgrades that strengthened security posture and operational efficiency.",
            "Supported mission-critical engineering programs with reliable compute access and responsive troubleshooting."
        ]
    )

    add_heading(document, "Earlier Career")
    early = document.add_paragraph()
    early.add_run(
        "Electronics procurement officer in the Israeli Navy and project manager at Elbit Systems, delivering complex electronic systems on tight schedules."
    )

    add_heading(document, "Education")
    edu = document.add_paragraph()
    edu.add_run("B.Sc., Electronic Engineering | Jerusalem College of Technology | 1991 – 1995")

    document.save("C:/Users/ryair/OneDrive - Intel Corporation/Documents/My paper - improved.docx")


if __name__ == "__main__":
    main()
